from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from docx import Document
from docx.shared import Inches
import io

class ExportService:
    @staticmethod
    def export_txt(summary_text, stats=None):
        content = "AI TEXT SUMMARIZATION REPORT\n"
        content += "=" * 50 + "\n\n"
        if stats:
            content += "SUMMARY STATISTICS:\n"
            content += f"Original Word Count: {stats.get('original_words', 'N/A')}\n"
            content += f"Summary Word Count: {stats.get('summary_words', 'N/A')}\n"
            content += f"Compression Ratio: {stats.get('compression_ratio', 'N/A')}\n"
            content += f"Method: {stats.get('method', 'N/A')}\n"
            content += f"Model Used: {stats.get('model_used', 'N/A')}\n"
            content += f"Processing Time: {stats.get('processing_time', 'N/A')}s\n\n"
        content += "SUMMARY:\n"
        content += "-" * 50 + "\n"
        content += summary_text + "\n"
        return content
    
    @staticmethod
    def export_pdf(summary_text, stats=None):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
        )
        story.append(Paragraph("AI Text Summarization Report", title_style))
        story.append(Spacer(1, 12))
        
        if stats:
            story.append(Paragraph("Summary Statistics", styles['Heading2']))
            story.append(Spacer(1, 12))
            stats_text = f"""
            <b>Original Word Count:</b> {stats.get('original_words', 'N/A')}<br/>
            <b>Summary Word Count:</b> {stats.get('summary_words', 'N/A')}<br/>
            <b>Compression Ratio:</b> {stats.get('compression_ratio', 'N/A')}<br/>
            <b>Method:</b> {stats.get('method', 'N/A')}<br/>
            <b>Model Used:</b> {stats.get('model_used', 'N/A')}<br/>
            <b>Processing Time:</b> {stats.get('processing_time', 'N/A')}s<br/>
            """
            story.append(Paragraph(stats_text, styles['Normal']))
            story.append(Spacer(1, 12))
        
        story.append(Paragraph("Summary", styles['Heading2']))
        story.append(Spacer(1, 12))
        story.append(Paragraph(summary_text.replace('\n', '<br/>'), styles['Normal']))
        
        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()
    
    @staticmethod
    def export_docx(summary_text, stats=None):
        doc = Document()
        doc.add_heading('AI Text Summarization Report', 0)
        
        if stats:
            doc.add_heading('Summary Statistics', level=1)
            doc.add_paragraph(f"Original Word Count: {stats.get('original_words', 'N/A')}")
            doc.add_paragraph(f"Summary Word Count: {stats.get('summary_words', 'N/A')}")
            doc.add_paragraph(f"Compression Ratio: {stats.get('compression_ratio', 'N/A')}")
            doc.add_paragraph(f"Method: {stats.get('method', 'N/A')}")
            doc.add_paragraph(f"Model Used: {stats.get('model_used', 'N/A')}")
            doc.add_paragraph(f"Processing Time: {stats.get('processing_time', 'N/A')}s")
        
        doc.add_heading('Summary', level=1)
        doc.add_paragraph(summary_text)
        
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
